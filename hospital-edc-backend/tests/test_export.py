# -*- coding: utf-8 -*-
"""导出 + 加密工具 + 报表"""
import io
import zipfile

import openpyxl
from sqlalchemy import text

from app.utils.crypto import encrypt, decrypt, mask
from app.services.scoring import calc_phq9_level
from tests.conftest import auth_header


def _create_patient(client, token, tag, center_id=None):
    payload = {
        "name_initials": tag, "full_name": f"导出{tag}", "id_card": "120103199001011234",
        "phone": "13800002222", "gender": "male", "age": 66,
        "enrollment_date": "2026-09-01",
    }
    if center_id:
        payload["center_id"] = center_id
    r = client.post("/api/patients/", headers=auth_header(token), json=payload)
    assert r.status_code == 200, r.text
    return r.json()


def test_crypto_roundtrip():
    """Fernet 加解密往返 + 历史明文透传 + 掩码"""
    cipher = encrypt("张三")
    assert cipher != "张三"
    assert decrypt(cipher) == "张三"
    assert decrypt("历史明文") == "历史明文"  # 非 token 原样返回
    assert mask("张三丰", 1) == "张**"
    assert mask("120103199001011234", 3) == "120****"
    assert mask("", 1) == ""

def test_scoring_levels():
    assert calc_phq9_level(3) is not None
    assert calc_phq9_level(0) != calc_phq9_level(20)

def test_export_xlsx_sheets_and_scores(client, admin_token):
    p = _create_patient(client, admin_token, "E1")
    r = client.post("/api/visits/", headers=auth_header(admin_token), json={
        "patient_id": p["id"], "visit_type": "baseline", "visit_date": "2026-09-02",
    })
    vid = r.json()["id"]
    client.post(f"/api/visits/{vid}/questionnaire", headers=auth_header(admin_token), json={
        "questionnaire_type": "phq9",
        "q1": 1, "q2": 1, "q3": 1, "q4": 1, "q5": 1, "q6": 1, "q7": 1, "q8": 1, "q9": 1,
    })

    r = client.get("/api/export/?format=xlsx&deidentify=true", headers=auth_header(admin_token))
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["患者主表", "访视宽表", "用药表", "不良事件表", "数据字典"]

    # 宽表含量表总分列
    header = [c.value for c in wb["访视宽表"][1]]
    for col in ("phq9_total", "phq9_level", "eq5d_vas", "dtsq_total"):
        assert col in header, f"缺少列 {col}"

    # 脱敏：主表姓名/身份证带掩码
    ws = wb["患者主表"]
    h = [c.value for c in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    row = dict(zip(h, rows[0]))
    assert "*" in row["full_name"]
    assert "*" in row["id_card"]

    # 字典覆盖宽表所有列
    dict_cols = {r[1] for r in wb["数据字典"].iter_rows(min_row=2, values_only=True)
                 if r[0] == "访视宽表"}
    missing = [c for c in header if c not in dict_cols]
    assert not missing, f"数据字典缺失: {missing}"

def test_export_privacy_requires_admin(client, researcher1_token):
    """研究者导含隐私数据应 403"""
    r = client.get("/api/export/?format=xlsx&deidentify=false",
                   headers=auth_header(researcher1_token))
    assert r.status_code == 403

def test_export_center_isolation(client, admin_token, researcher1_token, researcher2_token, seed):
    """分中心导出只含本中心患者"""
    _create_patient(client, researcher2_token, "E2")
    r = client.get("/api/export/?format=xlsx&deidentify=true",
                   headers=auth_header(researcher1_token))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["患者主表"]
    codes = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
    assert codes and all(str(c).startswith("TJ01-") for c in codes)

def test_export_csv_zip(client, admin_token):
    r = client.get("/api/export/?format=csv&deidentify=true", headers=auth_header(admin_token))
    assert r.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    assert zf.namelist() == ["患者主表.csv", "访视宽表.csv", "用药表.csv", "不良事件表.csv", "数据字典.csv"]
    text = zf.read("患者主表.csv").decode("utf-8-sig")
    assert "patient_code" in text

def test_export_audit_logged(client, admin_token, db):
    """每次导出写审计日志"""
    before = db.execute(text("SELECT COUNT(*) FROM audit_logs WHERE action='export'")).fetchone()[0]
    client.get("/api/export/?format=xlsx&deidentify=true", headers=auth_header(admin_token))
    after = db.execute(text("SELECT COUNT(*) FROM audit_logs WHERE action='export'")).fetchone()[0]
    assert after == before + 1

def test_reports_summary_no_patient_detail(client, researcher1_token):
    """分中心可看全局聚合，但接口不返回患者明细"""
    r = client.get("/api/reports/summary", headers=auth_header(researcher1_token))
    assert r.status_code == 200
    body = r.json()
    assert "patients_total" in body and "patients_by_center" in body
    assert "items" not in body and "patients" not in body

def test_reports_enrollment_isolated(client, researcher2_token):
    """分中心入组报表只统计本中心"""
    r = client.get("/api/reports/enrollment", headers=auth_header(researcher2_token))
    assert r.status_code == 200
    for c in r.json()["by_center"]:
        assert c["center_code"] == "TJ-02"
