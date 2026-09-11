# -*- coding: utf-8 -*-
"""M3 验收测试：数据导出（多 sheet xlsx / csv zip / 脱敏 / 中心隔离 / 导出审计）"""
import io
import json
import sys
import urllib.parse
import urllib.request
import zipfile

BASE = "http://127.0.0.1:8000"


def call(method, path, body=None, token=None, raw=False):
    req = urllib.request.Request(BASE + path, method=method)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    data = json.dumps(body).encode() if body is not None else None
    try:
        with urllib.request.urlopen(req, data) as r:
            content = r.read()
            if raw:
                return r.status, content
            return r.status, json.loads(content) if content else {}
    except urllib.error.HTTPError as e:
        content = e.read()
        try:
            return e.code, json.loads(content)
        except Exception:
            return e.code, content


def login(username, password):
    form = urllib.parse.urlencode({"username": username, "password": password}).encode()
    req = urllib.request.Request(BASE + "/api/auth/login", data=form, method="POST")
    req.add_header("Content-Type", "application/x-www-form-urlencoded")
    with urllib.request.urlopen(req) as r:
        data = json.loads(r.read())
    return data["access_token"]


admin = login("admin", "Admin@123")
researcher = login("researcher1", "Test@123")
print("admin/researcher login ok")

# ── 1. 解锁 visit 1，补 PHQ-9/GAD-7 数据，再锁定 ──
st, v = call("GET", "/api/visits/1", token=admin)
print("visit1:", st, v.get("status"), v.get("visit_type"))

if v.get("status") == "locked":
    st, r = call("POST", "/api/visits/1/unlock", token=admin)
    print("unlock:", st, r)
    assert st == 200

st, r = call("POST", "/api/visits/1/questionnaire", token=admin, body={
    "questionnaire_type": "phq9",
    "q1": 2, "q2": 1, "q3": 0, "q4": 1, "q5": 0, "q6": 2, "q7": 1, "q8": 0, "q9": 1,
})
print("phq9 save:", st, r if st != 200 else "ok")
assert st == 200 and r.get("total_score") == 8, r

st, r = call("POST", "/api/visits/1/questionnaire", token=admin, body={
    "questionnaire_type": "gad7",
    "q1": 1, "q2": 2, "q3": 0, "q4": 1, "q5": 0, "q6": 1, "q7": 2,
})
print("gad7 save:", st, r if st != 200 else "ok")
assert st == 200 and r.get("total_score") == 7, r

st, r = call("POST", "/api/visits/1/lock", token=admin)
print("re-lock:", st, r)
assert st == 200

# ── 2. 管理员导出 xlsx（脱敏）──
st, content = call("GET", "/api/export/?format=xlsx&deidentify=true", token=admin, raw=True)
print("export xlsx:", st, f"{len(content)} bytes")
assert st == 200, content

import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(content))
sheets = wb.sheetnames
print("sheets:", sheets)
assert sheets == ["患者主表", "访视宽表", "用药表", "不良事件表", "数据字典"], sheets

# 访视宽表列检查
ws = wb["访视宽表"]
header = [c.value for c in ws[1]]
print("visit wide header:", header)
for col in ["patient_code", "phq9_total", "phq9_level", "gad7_total", "gad7_level",
            "eq5d_vas", "dtsq_total", "sbp_mmhg", "hba1c"]:
    assert col in header, f"missing column {col}"
rows = list(ws.iter_rows(min_row=2, values_only=True))
print("visit rows:", len(rows))
by_code = {}
for row in rows:
    d = dict(zip(header, row))
    by_code.setdefault(d["patient_code"], []).append(d)
# visit1 的 phq9 总分应为 8
v1 = [d for d in by_code.get("TJ01-0001", []) if d["visit_type"] == "baseline"]
assert v1 and v1[0]["phq9_total"] == 8, f"phq9_total={v1[0]['phq9_total'] if v1 else None}"
print("phq9_total=8 ok, level:", v1[0]["phq9_level"], "gad7_total:", v1[0]["gad7_total"], "gad7_level:", v1[0]["gad7_level"])

# 患者主表脱敏检查
ws = wb["患者主表"]
header = [c.value for c in ws[1]]
row = dict(zip(header, list(ws.iter_rows(min_row=2, values_only=True))[0]))
print("patient row:", {k: row[k] for k in ("patient_code", "full_name", "id_card", "phone", "gender", "status")})
assert "*" in row["full_name"] and len(row["full_name"]) <= 5, "full_name 未脱敏"
assert "****" in row["id_card"], "id_card 未脱敏"

# 数据字典与访视宽表列一致性
ws = wb["数据字典"]
dict_rows = list(ws.iter_rows(values_only=True))
dict_cols = {r[1] for r in dict_rows[1:] if r[0] == "访视宽表"}
missing = [c for c in header if c not in dict_cols and c in [x for x in dict_cols]]  # 宽表列应在字典中
wide_cols = [c.value for c in wb["访视宽表"][1]]
not_in_dict = [c for c in wide_cols if c not in dict_cols]
print("宽表列不在数据字典中:", not_in_dict)
assert not not_in_dict, f"字典缺失: {not_in_dict}"

# ── 3. 管理员导出 xlsx（含隐私）──
st, content2 = call("GET", "/api/export/?format=xlsx&deidentify=false", token=admin, raw=True)
print("export xlsx privacy:", st, f"{len(content2)} bytes")
assert st == 200
wb2 = openpyxl.load_workbook(io.BytesIO(content2))
ws = wb2["患者主表"]
header = [c.value for c in ws[1]]
row = dict(zip(header, list(ws.iter_rows(min_row=2, values_only=True))[0]))
print("privacy patient row:", {k: row[k] for k in ("patient_code", "full_name", "id_card", "phone")})
assert "*" not in str(row["full_name"]), "full_name 不应脱敏"
assert row["full_name"], "full_name 为空（解密失败？）"

# ── 4. csv zip 导出 ──
st, content3 = call("GET", "/api/export/?format=csv&deidentify=true", token=admin, raw=True)
print("export csv:", st, f"{len(content3)} bytes")
assert st == 200
zf = zipfile.ZipFile(io.BytesIO(content3))
names = zf.namelist()
print("zip files:", names)
assert names == ["患者主表.csv", "访视宽表.csv", "用药表.csv", "不良事件表.csv", "数据字典.csv"], names
csv_text = zf.read("患者主表.csv").decode("utf-8-sig")
print("csv first line:", csv_text.splitlines()[0][:80])

# ── 5. 研究者权限：deidentify=false 应 403；只能导本中心 ──
st, r = call("GET", "/api/export/?format=xlsx&deidentify=false", token=researcher, raw=True)
print("researcher deidentify=false:", st)
assert st == 403, f"应 403，实际 {st}"

st, content4 = call("GET", "/api/export/?format=xlsx&deidentify=true", token=researcher, raw=True)
print("researcher export:", st, f"{len(content4)} bytes")
assert st == 200
wb3 = openpyxl.load_workbook(io.BytesIO(content4))
ws = wb3["患者主表"]
codes = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
centers = set(str(c).split("-")[0] for c in codes)
print("researcher sees patient codes:", codes, "centers:", centers)
assert all(str(c).startswith("TJ01") for c in codes), "研究者应只能看到本中心（TJ01）患者"

# ── 6. 导出审计 ──
st, logs = call("GET", "/api/audit-logs?limit=10", token=admin)
print("audit logs status:", st)
if st == 200:
    items = logs if isinstance(logs, list) else logs.get("items", [])
    exports = [l for l in items if l.get("action") == "export"]
    for e in exports[:3]:
        print("audit export:", e.get("table_name"), e.get("detail") or e.get("new_value"))

print("\n✅ M3 导出验收全部通过")
